"""成品导出：Markdown / TXT / DOCX 共用章节范围。"""
from io import BytesIO
import re

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.application.delivery import resolve_scope
from app.infrastructure import repository as repos


def render_markdown(title: str, chapters: list) -> str:
    parts = [f"# {title}\n"]
    for c in chapters:
        parts.append(f"## 第{c.seq}章 {c.title}\n\n{c.content}\n")
    return "\n".join(parts)


def render_txt(title: str, chapters: list) -> str:
    parts = [title, ""]
    for c in chapters:
        parts.append(f"第{c.seq}章 {c.title}\n{c.content}\n")
    return "\n".join(parts)


def render_docx(title: str, chapters: list) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=0)
    for index, chapter in enumerate(chapters):
        if index:
            document.add_page_break()
        document.add_heading(f"第{chapter.seq}章 {chapter.title}", level=1)
        for paragraph in re.split(r"\n\s*\n", chapter.content.strip()):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def export_file(db: Session, novel_id: str, request) -> bytes:
    novel = repos.NovelRepo(db).get(novel_id)
    _, chapters, _ = resolve_scope(db, novel_id, request)
    if not chapters:
        raise HTTPException(status_code=400, detail="所选范围没有可导出的章节")
    if request.format == "docx":
        return render_docx(novel.title, chapters)
    if request.format == "txt":
        return render_txt(novel.title, chapters).encode("utf-8")
    return render_markdown(novel.title, chapters).encode("utf-8")


def export_markdown(db: Session, novel_id: str) -> str:
    novel = repos.NovelRepo(db).get(novel_id)
    _, chapters, _ = resolve_scope(db, novel_id)
    return render_markdown(novel.title, chapters)


def export_txt(db: Session, novel_id: str) -> str:
    novel = repos.NovelRepo(db).get(novel_id)
    _, chapters, _ = resolve_scope(db, novel_id)
    return render_txt(novel.title, chapters)
