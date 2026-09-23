"""阶段八成品交付：统计、范围、预检查与三种文件格式。"""
from io import BytesIO
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.infrastructure import models
from app.infrastructure import repository as repos
from app.infrastructure.db import get_db
from app.main import app


engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
models.Base.metadata.create_all(bind=engine)
Session = sessionmaker(bind=engine, autocommit=False, autoflush=False)


def override_get_db():
    db = Session()
    try:
        yield db
    finally:
        db.close()


client = TestClient(app)


class DeliveryTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.dependency_overrides[get_db] = override_get_db

    def setUp(self):
        self.db = Session()
        self.novel_id = repos.NovelRepo(self.db).create(title="交付测试书", genre="都市").id
        outline = repos.OutlineRepo(self.db)
        self.volume = outline.create(self.novel_id, title="上卷", level="volume")
        act = outline.create(self.novel_id, title="第一幕", level="act", parent_id=self.volume.id)
        node = outline.create(self.novel_id, title="第一章节点", level="chapter", parent_id=act.id)
        self.second_volume = outline.create(self.novel_id, title="下卷", level="volume")
        chapters = repos.ChapterRepo(self.db)
        self.first = chapters.create(
            self.novel_id, seq=1, title="开篇", content="吴迪 出发。\n\n第二段。",
            status="finalized", outline_node_id=node.id, tension_score=5.5, drift_score=0.2,
        )
        self.draft = chapters.create(
            self.novel_id, seq=2, title="草稿", content="草稿 正文", status="draft",
            outline_node_id=node.id,
        )
        self.unbound = chapters.create(
            self.novel_id, seq=3, title="无大纲", content="第三章 正文", status="finalized",
        )
        self.last = chapters.create(
            self.novel_id, seq=4, title="结尾", content="结尾 正文", status="finalized",
            outline_node_id=self.second_volume.id,
        )

    def tearDown(self):
        self.db.close()

    def endpoint(self, suffix):
        return f"/api/v1/novels/{self.novel_id}/{suffix}"

    def post_file(self, **kwargs):
        return client.post(self.endpoint("export/file"), json=kwargs)

    def test_summary_counts_and_volume_ancestry(self):
        response = client.get(self.endpoint("book/summary"))
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["chapter_count"], 4)
        self.assertEqual(body["draft_count"], 1)
        self.assertEqual(body["finalized_count"], 3)
        self.assertEqual(body["word_count"], sum(c["word_count"] for c in body["chapters"]))
        self.assertEqual(body["chapters"][0]["volume_id"], self.volume.id)
        self.assertEqual(body["chapters"][0]["tension_score"], 5.5)
        self.assertEqual(body["chapters"][2]["volume_id"], "")

    def test_all_export_excludes_drafts_and_keeps_unbound(self):
        response = self.post_file(format="markdown", scope="all")
        self.assertEqual(response.status_code, 200, response.text)
        content = response.content.decode("utf-8")
        self.assertLess(content.index("第1章"), content.index("第3章"))
        self.assertLess(content.index("第3章"), content.index("第4章"))
        self.assertNotIn("第2章", content)
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])

    def test_volume_export_uses_ancestor_and_draft_toggle(self):
        request = {"format": "txt", "scope": "volume", "volume_id": self.volume.id}
        response = self.post_file(**request)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("第1章", response.text)
        self.assertNotIn("第2章", response.text)
        self.assertNotIn("第3章", response.text)
        included = self.post_file(**request, include_drafts=True)
        self.assertEqual(included.status_code, 200, included.text)
        self.assertIn("第2章", included.text)

    def test_selected_chapters_order_and_cross_novel_rejection(self):
        response = self.post_file(
            scope="chapters", chapter_ids=[self.last.id, self.first.id], format="markdown"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertLess(response.text.index("第1章"), response.text.index("第4章"))
        self.assertNotIn("第3章", response.text)
        other_id = repos.NovelRepo(self.db).create(title="另一作品", genre="都市").id
        foreign = repos.ChapterRepo(self.db).create(other_id, title="私有章", content="私有正文")
        bad = self.post_file(scope="chapters", chapter_ids=[self.first.id, foreign.id])
        self.assertEqual(bad.status_code, 400)

    def test_invalid_or_empty_scope_returns_400(self):
        self.assertEqual(self.post_file(scope="volume", volume_id="unknown").status_code, 400)
        self.assertEqual(self.post_file(scope="chapters", chapter_ids=[]).status_code, 400)
        self.assertEqual(
            self.post_file(scope="chapters", chapter_ids=[self.draft.id]).status_code, 400
        )
        self.assertEqual(
            self.post_file(scope="volume", volume_id=self.second_volume.id, format="docx").status_code,
            200,
        )

    def test_preflight_reports_drafts_unbound_foreshadowing_and_rules_without_llm(self):
        repos.ForeshadowingRepo(self.db).create(
            self.novel_id, content="暗号", status="planted"
        )
        repos.ForeshadowingRepo(self.db).create(
            self.novel_id, content="吴迪", status="consumed"
        )
        with patch("app.application.consistency.llm.complete", side_effect=AssertionError("LLM called")):
            response = client.post(self.endpoint("export/preflight"), json={"scope": "all"})
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["chapter_count"], 3)
        self.assertEqual([c["id"] for c in body["draft_chapters"]], [self.draft.id])
        self.assertEqual([c["id"] for c in body["unbound_chapters"]], [self.unbound.id])
        self.assertEqual(len(body["unresolved_foreshadowings"]), 1)
        self.assertEqual(body["rule_conflicts"][0]["chapter_id"], self.first.id)
        self.assertFalse(body["ready"])

    def test_preflight_unresolved_foreshadowing_is_warning_only(self):
        repos.ForeshadowingRepo(self.db).create(self.novel_id, content="悬念", status="planted")
        response = client.post(
            self.endpoint("export/preflight"),
            json={"scope": "chapters", "chapter_ids": [self.last.id]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["ready"])
        self.assertEqual(len(response.json()["unresolved_foreshadowings"]), 1)

    def test_docx_structure_and_mime(self):
        from docx import Document

        response = self.post_file(format="docx", scope="all")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        document = Document(BytesIO(response.content))
        headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, ["第1章 开篇", "第3章 无大纲", "第4章 结尾"])
        paragraphs = [p.text for p in document.paragraphs]
        self.assertIn("吴迪 出发。", paragraphs)
        self.assertIn("第二段。", paragraphs)
        self.assertEqual(paragraphs[0], "交付测试书")
        self.assertGreaterEqual(
            sum('w:type="page"' in p._p.xml for p in document.paragraphs), 2
        )

    def test_legacy_get_exports_remain_finalized_whole_book(self):
        for suffix in ("markdown", "txt"):
            response = client.get(self.endpoint(f"export/{suffix}"))
            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn("第1章", response.text)
            self.assertIn("第3章", response.text)
            self.assertNotIn("第2章", response.text)

    def test_three_formats_share_identical_scope_selection(self):
        from docx import Document

        cases = (
            ("all", {}, [1, 3, 4]),
            ("volume", {"volume_id": self.volume.id}, [1]),
            ("chapters", {"chapter_ids": [self.last.id, self.first.id]}, [1, 4]),
        )
        for scope, selection, expected in cases:
            for file_format in ("markdown", "txt", "docx"):
                with self.subTest(scope=scope, format=file_format):
                    response = self.post_file(scope=scope, format=file_format, **selection)
                    self.assertEqual(response.status_code, 200, response.text)
                    if file_format == "docx":
                        document = Document(BytesIO(response.content))
                        text = "\n".join(p.text for p in document.paragraphs)
                    else:
                        text = response.content.decode("utf-8")
                    observed = [seq for seq in range(1, 5) if f"第{seq}章" in text]
                    self.assertEqual(observed, expected)
                    self.assertEqual(
                        [text.index(f"第{seq}章") for seq in expected],
                        sorted(text.index(f"第{seq}章") for seq in expected),
                    )


if __name__ == "__main__":
    unittest.main()
